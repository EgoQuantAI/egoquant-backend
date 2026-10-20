"""Merge the four returned labeling files. Needs openpyxl.

    python eval/merge_labels.py                      reads eval/labeling/person1.xlsx ... person4.xlsx
    python eval/merge_labels.py --disagreements eval/labeling/disagreements_labeled.xlsx
                                                     use a filled adjudication file with another name

Checks that every row has a valid label (positive / negative / neutral), that each file holds exactly the ids
of its two blocks in eval/labeling_sheet.csv, and that every headline was labeled by exactly two people.
Writes eval/labels_merged.csv (id, headline, label_a, label_b, agree; gitignored) and prints per-person label
counts, the agreement rate, Cohen's kappa for each pair of labelers sharing a block, the disagreement types
and every disagreement. Nothing is written if any check fails (exit code 1).

Obvious misspellings (TYPO_FIXES) are corrected and every correction is printed; the returned files are never
modified. Short rows (e.g. files saved from Google Sheets) are padded.

Adjudication (second run):
  * If there are disagreements and eval/labeling/disagreements.xlsx does not exist yet, it is created: one row
    per disagreement (headline, fund, the two labels without names, an empty "final" column with a dropdown)
    plus an Instructions sheet. An existing file is never overwritten (delete it to rebuild it).
  * Once every "final" cell is filled, eval/headlines_labeled.csv (id, headline, label, labeler) is written:
    agreed rows get labeler "agreed", resolved rows "adjudicated". It is gitignored: labels are never committed.
"""

from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "eval"))

from make_labeling_handouts import fund_text  # noqa: E402
from make_labeling_sheet import ASSIGNMENT  # noqa: E402

LABELS = {"positive", "negative", "neutral"}
# Obvious misspellings only; anything else is reported as invalid.
TYPO_FIXES = {
    "netural": "neutral", "nuetral": "neutral", "neutal": "neutral", "neural": "neutral", "nutral": "neutral",
    "postive": "positive", "positve": "positive", "possitive": "positive", "posiitve": "positive",
    "negtive": "negative", "negaitve": "negative", "negatve": "negative", "negativ": "negative", "neagtive": "negative",
}
HEADERS = ["row", "headline", "fund", "label", "id (do not edit)"]
MERGED_COLUMNS = ["id", "headline", "label_a", "label_b", "agree"]
LABELED_COLUMNS = ["id", "headline", "label", "labeler"]
DIS_HEADERS = ["row", "headline", "fund", "label 1", "label 2", "final", "id (do not edit)"]
DIS_INSTRUCTIONS = [
    ("Resolve the disagreements", True),
    ("Two people labeled each of these headlines differently. Agree on one final label together and choose it", False),
    ("in the 'final' column from the dropdown: positive, negative or neutral.", False),
    ("", False),
    ("Judge for an investor who holds the listed fund:", False),
    ("    positive  =  likely good news for that investor", False),
    ("    negative  =  likely bad news for that investor", False),
    ("    neutral   =  factual, mixed or unclear for that investor", False),
    ("Judge the headline only. Do not open links or read the article.", False),
    ("If you are still unsure, choose neutral.", False),
    ("No AI tools. Only fill the 'final' column; do not edit, sort or delete rows.", False),
]


def normalize(value) -> tuple[str, str | None]:
    """(label, fix note or None). Strips spaces, lower-cases, corrects TYPO_FIXES."""
    label = str(value or "").strip().lower()
    if label in TYPO_FIXES:
        return TYPO_FIXES[label], f"{value!r} -> {TYPO_FIXES[label]!r}"
    return label, None


FIXES: list[str] = []                                 # corrections made while reading (printed by main)


def read_labels(path: Path) -> tuple[list[tuple[str, str]], list[str]]:
    """[(id, label)] from a returned workbook, plus problems found while reading."""
    from openpyxl import load_workbook
    problems: list[str] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        return [], [f"{path.name}: file not found"]
    if "Labels" not in wb.sheetnames:
        return [], [f"{path.name}: sheet 'Labels' is missing"]
    rows = [tuple(r) + (None,) * max(0, 5 - len(r)) for r in wb["Labels"].iter_rows(values_only=True)]
    if not rows or [str(h or "").strip() for h in rows[0][:5]] != HEADERS:
        return [], [f"{path.name}: unexpected header row {rows[0] if rows else None}"]
    out = []
    for n, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in row):
            continue                                    # trailing empty rows
        rid = str(row[4] or "").strip()
        label, fix = normalize(row[3])
        if fix:
            FIXES.append(f"{path.name} row {n} (id {rid}): {fix}")
        if label not in LABELS:
            problems.append(f"{path.name} row {n}: label {row[3]!r} is not positive, negative or neutral")
        out.append((rid, label))
    return out, problems


def merge(sheet: list[dict], files: dict[str, Path]) -> tuple[list[dict], list[str]]:
    """(merged rows in sheet order, problems). Merged rows are only valid when problems is empty."""
    by_id = {r["id"]: r for r in sheet}
    problems: list[str] = []
    labels: dict[str, list[tuple[int, str]]] = defaultdict(list)
    for person, path in files.items():
        expected = {r["id"] for r in sheet if r["block"] in ASSIGNMENT[person]}
        rows, read_problems = read_labels(path)
        problems += read_problems
        if not rows and read_problems:
            continue
        ids = [rid for rid, _ in rows]
        dup = sorted({i for i in ids if ids.count(i) > 1})
        missing, extra = sorted(expected - set(ids)), sorted(set(ids) - expected)
        if dup:
            problems.append(f"{path.name}: duplicate ids {dup}")
        if missing:
            problems.append(f"{path.name}: {len(missing)} ids of blocks {'+'.join(ASSIGNMENT[person])} missing: {missing[:5]}")
        if extra:
            problems.append(f"{path.name}: ids not in its blocks (edited?): {extra[:5]}")
        for rid, label in rows:
            if rid in expected:
                labels[rid].append((int(person.removeprefix("person")), label))
    for rid in by_id:
        if len(labels[rid]) != 2 and not problems:
            problems.append(f"id {rid}: labeled {len(labels[rid])} times, expected 2")
    merged = []
    for r in sheet:
        pair = [lab for _, lab in sorted(labels[r["id"]])]
        if len(pair) == 2:
            merged.append({"id": r["id"], "headline": r["headline"], "label_a": pair[0], "label_b": pair[1],
                           "agree": "yes" if pair[0] == pair[1] else "no"})
    return merged, problems


def cohen_kappa(a: list[str], b: list[str]) -> float | None:
    """Cohen's kappa for two labelers on the same items (None if undefined)."""
    n = len(a)
    if n == 0:
        return None
    observed = sum(x == y for x, y in zip(a, b)) / n
    expected = sum((a.count(c) / n) * (b.count(c) / n) for c in LABELS)
    return None if expected == 1 else (observed - expected) / (1 - expected)


def pair_kappas(sheet: list[dict], per_person: dict[str, dict[str, str]]) -> list[tuple[str, str, str, int, float | None]]:
    """(person_x, person_y, shared block, n, kappa) for every pair of labelers that shares a block."""
    out = []
    people = list(ASSIGNMENT)
    for i, x in enumerate(people):
        for y in people[i + 1:]:
            shared = sorted(set(ASSIGNMENT[x]) & set(ASSIGNMENT[y]))
            if not shared:
                continue
            ids = [r["id"] for r in sheet if r["block"] in shared]
            a = [per_person[x][rid] for rid in ids if rid in per_person[x] and rid in per_person[y]]
            b = [per_person[y][rid] for rid in ids if rid in per_person[x] and rid in per_person[y]]
            out.append((x, y, "+".join(shared), len(a), cohen_kappa(a, b)))
    return out


def write_disagreements(path: Path, disagreements: list[dict], sheet_by_id: dict[str, dict]) -> None:
    """The adjudication workbook: one row per disagreement, labels without names, a 'final' dropdown."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Disagreements"
    ws.append(DIS_HEADERS)
    for n, r in enumerate(disagreements, start=1):
        ws.append([n, r["headline"], fund_text(sheet_by_id[r["id"]]["matched_tickers"]), r["label_a"], r["label_b"],
                   None, r["id"]])
    base, bold = Font(size=12), Font(size=12, bold=True)
    grey_font, grey_fill = Font(size=11, color="808080"), PatternFill("solid", fgColor="EDEDED")
    for cell in ws[1]:
        cell.font, cell.fill = bold, PatternFill("solid", fgColor="D9E1F2")
    last = len(disagreements) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = base
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in ("B", "C"))
        row[6].font, row[6].fill = grey_font, grey_fill
    ws["G1"].fill = grey_fill
    for col, width in {"A": 6, "B": 75, "C": 40, "D": 11, "E": 11, "F": 13, "G": 20}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"positive,negative,neutral"', allow_blank=True, showErrorMessage=True,
                        errorStyle="stop", errorTitle="Invalid label",
                        error="Please choose positive, negative or neutral from the list.")
    ws.add_data_validation(dv)
    dv.add(f"F2:F{max(last, 2)}")
    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width = 110
    for text, heading in DIS_INSTRUCTIONS:
        info.append([text])
        info.cell(row=info.max_row, column=1).font = Font(size=13 if heading else 12, bold=heading)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_finals(path: Path, disagreements: list[dict]) -> tuple[dict[str, str], list[str]]:
    """{id: final label} from a filled disagreements.xlsx, plus problems (missing/invalid finals, wrong ids)."""
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)["Disagreements"]
    rows = [tuple(r) + (None,) * max(0, 7 - len(r)) for r in ws.iter_rows(values_only=True)]
    if [str(h or "").strip() for h in rows[0][:7]] != DIS_HEADERS:
        return {}, [f"{path.name}: unexpected header row {rows[0][:7]}"]
    finals, problems = {}, []
    for n, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in row):
            continue
        rid = str(row[6] or "").strip()
        label, fix = normalize(row[5])
        if fix:
            FIXES.append(f"{path.name} row {n} (id {rid}): {fix}")
        if label not in LABELS:
            problems.append(f"{path.name} row {n}: final {row[5]!r} is not positive, negative or neutral")
        finals[rid] = label
    expected = {r["id"] for r in disagreements}
    if set(finals) != expected:
        problems.append(f"{path.name}: ids do not match the current disagreements "
                        f"(missing {sorted(expected - set(finals))[:5]}, extra {sorted(set(finals) - expected)[:5]})")
    return finals, problems


def final_labels(merged: list[dict], finals: dict[str, str]) -> list[dict]:
    """Rows for headlines_labeled.csv: agreed rows -> 'agreed', resolved rows -> 'adjudicated'."""
    out = []
    for r in merged:
        if r["agree"] == "yes":
            out.append({"id": r["id"], "headline": r["headline"], "label": r["label_a"], "labeler": "agreed"})
        else:
            out.append({"id": r["id"], "headline": r["headline"], "label": finals[r["id"]], "labeler": "adjudicated"})
    return out


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=columns, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def main(argv: list[str] | None = None) -> int:
    import argparse
    ap = argparse.ArgumentParser(description="Merge the returned labeling files.")
    ap.add_argument("folder", nargs="?", default=None, help="folder with person1.xlsx ... (default eval/labeling)")
    ap.add_argument("--disagreements", default=None,
                    help="adjudication workbook to read (default: <folder>/disagreements.xlsx)")
    args = ap.parse_args(argv or [])
    folder = Path(args.folder) if args.folder else ROOT / "eval" / "labeling"
    dis_path = Path(args.disagreements).expanduser() if args.disagreements else folder / "disagreements.xlsx"
    with open(ROOT / "eval" / "labeling_sheet.csv", newline="", encoding="utf-8") as fh:
        sheet = list(csv.DictReader(fh))
    FIXES.clear()
    files = {p: folder / f"{p}.xlsx" for p in ASSIGNMENT}
    merged, problems = merge(sheet, files)
    for f in FIXES:
        print("fixed typo:", f)
    if problems:
        print("Not merged. Fix these first:")
        for p in problems:
            print("  -", p)
        return 1
    out = ROOT / "eval" / "labels_merged.csv"
    write_csv(out, merged, MERGED_COLUMNS)
    agree = sum(r["agree"] == "yes" for r in merged)
    print(f"wrote {len(merged)} rows to {out}")
    per_person = {p: dict(read_labels(path)[0]) for p, path in files.items()}
    print("\nlabels per person:")
    for p, labels in per_person.items():
        counts = {c: sum(v == c for v in labels.values()) for c in ("positive", "negative", "neutral")}
        print(f"  {p}: " + ", ".join(f"{c} {n}" for c, n in counts.items()) + f" (total {len(labels)})")
    print(f"\nagreement: {agree}/{len(merged)} = {agree / len(merged):.1%}")
    print("Cohen's kappa per pair (shared block):")
    for x, y, block, n, k in pair_kappas(sheet, per_person):
        print(f"  {x} + {y} (block {block}, {n} headlines): " + ("undefined" if k is None else f"{k:.3f}"))
    kinds: dict[str, int] = {}
    for r in merged:
        if r["agree"] == "no":
            key = " vs ".join(sorted((r["label_a"], r["label_b"])))
            kinds[key] = kinds.get(key, 0) + 1
    print("disagreement types:", ", ".join(f"{k} {v}" for k, v in sorted(kinds.items(), key=lambda kv: -kv[1])) or "none")
    disagreements = [r for r in merged if r["agree"] == "no"]
    print(f"\n{len(disagreements)} disagreements (resolve together; still unsure -> neutral):")
    for r in disagreements:
        print(f"  {r['id']}  {r['label_a']:>8} vs {r['label_b']:<8}  {r['headline']}")
    return adjudicate(merged, disagreements, sheet, dis_path)


def adjudicate(merged: list[dict], disagreements: list[dict], sheet: list[dict], dis_path: Path) -> int:
    labeled_path = ROOT / "eval" / "headlines_labeled.csv"
    finals: dict[str, str] = {}
    if disagreements:
        if not dis_path.exists():
            write_disagreements(dis_path, disagreements, {r["id"]: r for r in sheet})
            print(f"\ncreated {dis_path} ({len(disagreements)} rows). Fill the 'final' column, then run this again.")
            return 0
        finals, problems = read_finals(dis_path, disagreements)
        for f in FIXES:
            if f.startswith(dis_path.name):
                print("fixed typo:", f)
        if problems:
            print(f"\n{dis_path.name} is not complete yet; headlines_labeled.csv not written:")
            for p in problems:
                print("  -", p)
            return 1
    rows = final_labels(merged, finals)
    write_csv(labeled_path, rows, LABELED_COLUMNS)
    n_adj = sum(r["labeler"] == "adjudicated" for r in rows)
    print(f"\nwrote {len(rows)} labeled headlines to {labeled_path} ({len(rows) - n_adj} agreed, {n_adj} adjudicated)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
