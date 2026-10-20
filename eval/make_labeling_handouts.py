"""Build one Excel handout per labeler from eval/labeling_sheet.csv. Needs openpyxl.

    python eval/make_labeling_handouts.py

Writes eval/labeling/person1.xlsx ... person4.xlsx (the folder is gitignored; never commit these files).
Each file holds only that person's two blocks (60 rows), shuffled with a fixed seed per person so the
blocks are not visible, a label dropdown (positive / negative / neutral) and an Instructions sheet.
No scoring, no labels, no model.
"""

from __future__ import annotations

import csv
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "eval"))

from egoquant.briefing.fund_tone import FUND_PLAIN_NAMES, fund_text  # noqa: E402,F401
from make_labeling_sheet import ASSIGNMENT  # noqa: E402

LABELS = ("positive", "negative", "neutral")
OUT_DIR = ROOT / "eval" / "labeling"
BASE_SEED = 42                                   # person N uses seed 42 + N
DEADLINE = "Send this file back to Bartu by Monday 14:00 Abu Dhabi time."

# Plain-English fund names: one source of truth, shared with Agent D's fund-aware tone (same text the model saw).
PLAIN_NAMES = FUND_PLAIN_NAMES

INSTRUCTIONS = [
    ("How to label", True),
    ("For each headline, choose a label in the 'label' column from the dropdown: positive, negative or neutral.", False),
    ("", False),
    ("Imagine an investor who holds the fund listed next to the headline. Is the headline:", False),
    ("    positive  =  likely good news for that investor", False),
    ("    negative  =  likely bad news for that investor", False),
    ("    neutral   =  factual, mixed or unclear for that investor", False),
    ("If several funds are listed and the news points in different directions for them, choose neutral.", False),
    ("", False),
    ("Rules", True),
    ("1. Judge the headline only. Do not open links or read the article.", False),
    ("2. If you are unsure, choose neutral.", False),
    ("3. Label alone. Do not discuss the headlines with anyone until everyone has finished.", False),
    ("4. No AI tools. Do not use ChatGPT, Claude or any other tool to suggest a label. Your own judgment only.", False),
    ("5. Only fill the 'label' column. Do not change, add, delete or sort rows, and do not edit the grey id column.", False),
    ("", False),
    ("Example", True),
    ("Headline: \"Treasury yields jump\"", False),
    ("    Fund TLT - long-term US Treasury bonds  ->  negative (when yields rise, long bond prices fall)", False),
    ("    Fund BIL - 1-3 month US Treasury bills  ->  neutral (very short bills barely move, and new bills pay more)", False),
    ("The same headline can deserve different labels for different funds: always judge for the listed fund.", False),
    ("", False),
    (DEADLINE, True),
]


def person_rows(sheet: list[dict], person: str) -> list[dict]:
    blocks = set(ASSIGNMENT[person])
    rows = sorted((r for r in sheet if r["block"] in blocks), key=lambda r: r["id"])
    random.Random(BASE_SEED + int(person.removeprefix("person"))).shuffle(rows)
    return rows


def write_handout(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    headers = ["row", "headline", "fund", "label", "id (do not edit)"]
    ws.append(headers)
    for n, r in enumerate(rows, start=1):
        ws.append([n, r["headline"], fund_text(r["matched_tickers"]), None, r["id"]])

    base, bold = Font(size=12), Font(size=12, bold=True)
    grey_font, grey_fill = Font(size=11, color="808080"), PatternFill("solid", fgColor="EDEDED")
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(vertical="center")
    last = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font = base
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in ("B", "C"))
        row[4].font, row[4].fill = grey_font, grey_fill
    ws["E1"].fill = grey_fill
    for col, width in {"A": 6, "B": 80, "C": 42, "D": 14, "E": 20}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True,
                        showDropDown=False, showErrorMessage=True, errorStyle="stop",
                        errorTitle="Invalid label", error="Please choose positive, negative or neutral from the list.",
                        showInputMessage=True, promptTitle="Label", prompt="positive, negative or neutral")
    ws.add_data_validation(dv)
    dv.add(f"D2:D{last}")

    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width = 110
    for text, is_heading in INSTRUCTIONS:
        info.append([text])
        cell = info.cell(row=info.max_row, column=1)
        cell.font = Font(size=13, bold=True) if is_heading else Font(size=12)
        cell.alignment = Alignment(wrap_text=True)
    wb.active = 1                                   # open on the instructions first
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    with open(ROOT / "eval" / "labeling_sheet.csv", newline="", encoding="utf-8") as fh:
        sheet = list(csv.DictReader(fh))
    for person in ASSIGNMENT:
        rows = person_rows(sheet, person)
        path = OUT_DIR / f"{person}.xlsx"
        write_handout(path, rows)
        print(f"{path}: {len(rows)} rows (blocks {' + '.join(ASSIGNMENT[person])})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
